# -*- coding: utf-8 -*-
"""
Excel 用例解析器

将 Excel 文件格式的用例转换为与 YAML 完全相同的数据结构，
让 TestRunner 无需区分来源。

Excel 列规范（可自定义，以下为默认）:
    A: 一级模块
    B: 二级模块
    C: 用例标题
    D: 步骤描述
    E: 操作类型
    F: 元素定位名称（对应 context.yaml 的 _WEB页面元素 别名）
    G: 输入数据 / 预期值
    H: 备注

每行 = 一个步骤；相同用例标题的若干行 = 一个用例。
"""

from __future__ import annotations

import json
import os
import ast
from typing import Any, Dict, List

from HAT.core.globalContext import g_context

try:
    from openpyxl import load_workbook
    _HAS_XLSX = True
except ImportError:
    _HAS_XLSX = False


# ============================================================
# 从 Excel 加载全局配置
# ============================================================
def load_context_from_excel(context_file: str) -> None:
    """从 Excel 加载配置（可选功能）

    如果目录中存在 context.xlsx，其第一页可按 『变量名 | 变量值 JSON』
    的两列结构提供配置。大多数情况下，直接使用 context.yaml 更简洁，
    这里只做基本支持。
    """
    if not _HAS_XLSX or not os.path.isfile(context_file):
        return

    wb = load_workbook(context_file, data_only=True)
    ws = wb.active
    for row in ws.iter_rows(min_row=2, values_only=True):  # 首行视为表头
        if not row or row[0] is None:
            continue
        key = str(row[0]).strip()
        value = row[1] if len(row) > 1 else ""
        value = safe_convert_value(value)
        g_context().set_dict(key, value)


# ============================================================
# 智能类型转换
# ============================================================
def safe_convert_value(value: Any) -> Any:
    """将字符串智能转换为 Python 类型

    规则:
    - None / 空字符串 → ""
    - 已经是非字符串 → 直接返回
    - 以 { 或 [ 开头 → 尝试 JSON / ast.literal_eval 解析
    - "True" / "False" / "None" → 对应 Python 值
    - 可转为数字 → int / float
    - 其他 → 原字符串
    """
    if value is None:
        return ""
    if not isinstance(value, str):
        return value

    stripped = value.strip()
    if not stripped:
        return ""

    # 1. JSON 解析
    if stripped[0] in "{[":
        try:
            return json.loads(stripped)
        except Exception:
            pass
        try:
            return ast.literal_eval(stripped)
        except Exception:
            return value

    # 2. 布尔 / None
    if stripped.lower() == "true":
        return True
    if stripped.lower() == "false":
        return False
    if stripped.lower() == "none" or stripped.lower() == "null":
        return None

    # 3. 数字
    try:
        if "." in stripped:
            return float(stripped)
        return int(stripped)
    except ValueError:
        pass

    return value


# ============================================================
# 核心：扫描 Excel 文件并归并为用例
# ============================================================
def group_cases_by_title(rows: List[tuple]) -> Dict[str, Dict[str, Any]]:
    """按『用例标题』将行归并为用例结构

    Args:
        rows: 每一行的元组 (一级模块, 二级模块, 用例标题, 步骤描述,
             操作类型, 元素定位名称, 数据内容/参数, 备注)

    Returns:
        { 用例标题: { 基础配置, 用例步骤: [...] }, ... }
    """
    grouped: Dict[str, Dict[str, Any]] = {}
    ordered_titles: List[str] = []

    for row in rows:
        cols = list(row) + [None] * (8 - len(row))  # 保证 8 列
        module1, module2, title, step_desc, action, element, data, remark = cols

        if title is None or str(title).strip() == "":
            # 空标题的行忽略（可能是分隔或注释）
            continue

        title = str(title).strip()
        if title not in grouped:
            grouped[title] = {
                "基础配置": {
                    "用例类型": "WebCase",
                    "一级模块": str(module1 or "未分类").strip(),
                    "二级模块": str(module2 or "未分类").strip(),
                    "用例标题": title,
                },
                "用例步骤": [],
            }
            ordered_titles.append(title)

        # 构建一个步骤字典
        step_dict: Dict[str, Any] = {}
        if step_desc:
            step_dict["步骤描述"] = str(step_desc).strip()
        if action:
            step_dict["操作类型"] = str(action).strip()
        if element:
            step_dict["_页面元素"] = str(element).strip()
        if data is not None and str(data).strip() != "":
            data_conv = safe_convert_value(data)
            # 如果数据是 dict，把它的键合并到步骤参数
            if isinstance(data_conv, dict):
                for k, v in data_conv.items():
                    step_dict[str(k)] = v
            else:
                # 否则用一个统一的字段名
                step_dict["数据内容"] = data_conv

        if not step_dict.get("操作类型"):
            # 没有操作类型的行不构成为步骤
            continue

        grouped[title]["用例步骤"].append(step_dict)

    return grouped


def load_excel_files(cases_dir: str) -> Dict[str, Any]:
    """扫描目录读取符合规则的 Excel 文件

    Returns:
        {"case_infos": [...], "case_names": [...]}
    """
    if not _HAS_XLSX:
        raise RuntimeError("缺少 openpyxl 依赖，请执行: pip install openpyxl")

    if not os.path.isdir(cases_dir):
        raise NotADirectoryError(f"用例目录不存在: {cases_dir}")

    case_infos: List[Dict[str, Any]] = []
    case_names: List[str] = []

    for fname in os.listdir(cases_dir):
        if not (fname.lower().endswith(".xlsx") or fname.lower().endswith(".xls")):
            continue
        if fname.startswith("~$"):
            continue  # Excel 临时文件

        fpath = os.path.join(cases_dir, fname)
        try:
            wb = load_workbook(fpath, data_only=True)
        except Exception as e:
            print(f"  ⚠ 无法读取 Excel {fname}: {e}")
            continue

        ws = wb.active

        # 读取数据行（跳过第一行表头）
        rows: List[tuple] = []
        for idx, row in enumerate(ws.iter_rows(values_only=True)):
            if idx == 0:
                continue  # 表头
            if row is None or all(cell is None for cell in row):
                continue
            rows.append(tuple(row))

        grouped = group_cases_by_title(rows)
        for title, case in grouped.items():
            if not case.get("用例步骤"):
                continue
            case_infos.append(case)
            case_names.append(title)

    return {"case_infos": case_infos, "case_names": case_names}


def excel_case_parser(cases_dir: str) -> Dict[str, Any]:
    """Excel 用例解析的统一入口

    Args:
        cases_dir: 用例目录路径

    Returns:
        {"case_infos": [...], "case_names": [...]}
    """
    # 先检查是否有 context.yaml（跨格式共享配置）
    from HAT.parse.YamlCaseParser import load_context_from_yaml
    context_yaml = os.path.join(cases_dir, "context.yaml")
    load_context_from_yaml(context_yaml)

    # 再检查是否有 context.xlsx
    context_xlsx = os.path.join(cases_dir, "context.xlsx")
    load_context_from_excel(context_xlsx)

    return load_excel_files(cases_dir)
